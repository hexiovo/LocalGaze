"""Input discovery and format-preserving output helpers."""

from __future__ import annotations

import csv
from decimal import Decimal, InvalidOperation
import hashlib
import json
import math
import shutil
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from LocalGaze.capture_paths import (
    INTERNAL_LAYOUT,
    LEGACY_LAYOUT,
    existing_sidecar_path,
    normalize_sidecar_layout,
    owned_sidecar_path,
    sidecar_path,
)


DATA_SUFFIXES = {".xlsx", ".xlsm", ".csv", ".tsv", ".txt"}
MODEL_SUFFIXES = {".xlsx", ".xlsm", ".pkl"}

ALIASES = {
    "time": ("时间", "时间戳", "timestamp", "time", "datetime"),
    "x": ("x", "gazex", "gaze_x", "水平坐标", "横坐标"),
    "y": ("y", "gazey", "gaze_y", "垂直坐标", "纵坐标"),
    "blink": ("blink", "眨眼", "眨眼标记"),
    "left_pupil": ("左瞳孔直径", "leftpupil", "left_pupil", "pupil_left"),
    "right_pupil": ("右瞳孔直径", "rightpupil", "right_pupil", "pupil_right"),
}

ANALYSIS_ALIASES = {
    # Raw coordinate candidates deliberately exclude Pre_X / Pre_Y and other
    # display-smoothed or cleaned fields.
    "x_raw": ("analysis_x_raw", "gaze_x_raw", "raw_x", "Pre_XRaw", "XRaw", "gaze_x", "gaze_x_px", "gazex", "x_px", "水平坐标", "横坐标", "x"),
    "y_raw": ("analysis_y_raw", "gaze_y_raw", "raw_y", "Pre_YRaw", "YRaw", "gaze_y", "gaze_y_px", "gazey", "y_px", "垂直坐标", "纵坐标", "y"),
    "received_monotonic_ns": ("received_monotonic_ns", "monotonic_ns"),
    "elapsed_s": ("elapsed_s", "elapsed_seconds", "Pre_TimeSeconds"),
    "time": ("time", "timestamp", "datetime", "时间", "时间戳", "time_s", "timestamp_s", "time_ms", "timestamp_ms", "time_us", "timestamp_us", "时间秒", "时间毫秒", "时间微秒"),
    "session_id": ("session_id", "session", "会话ID"),
    "sample_status": ("analysis_sample_status", "sample_status", "status", "状态"),
    "sample_accepted": ("sample_accepted", "accepted"),
    "gaze_valid": ("analysis_gaze_valid", "Pre_ValidMeasured", "ValidMeasured", "gaze_valid"),
    "blink": ("blink", "眨眼", "眨眼标记"),
    "interpolated": ("analysis_interpolated", "Pre_InterpolatedMask", "InterpolatedMask"),
    "bad_segment": ("analysis_bad_segment", "Pre_BadSegmentMask", "BadSegmentMask"),
    "break_before": ("analysis_break_before", "Pre_BreakBefore", "BreakBefore"),
    "segment_id": ("analysis_segment_id", "Pre_SegmentID", "SegmentID", "segment_id"),
    "trial_id": ("analysis_trial_id", "trial_id", "trial", "试次ID", "试次"),
    "stimulus_id": ("analysis_stimulus_id", "stimulus_id", "stimulus", "刺激ID", "刺激"),
    "invalid_reason": ("analysis_invalid_reason", "sample_status_reason", "invalid_reason"),
    "clock_quality": ("analysis_clock_quality", "clock_quality"),
    "time_source": ("analysis_time_source", "time_source"),
}

ANALYSIS_STATUS_VALUES = {
    "accepted", "no_face", "blink", "warmup", "invalid",
    "duplicate_frame_id", "processing_error",
}


def discover_files(path: Path, suffixes: set[str]) -> list[Path]:
    if path.is_file():
        return [path]
    if not path.is_dir():
        raise FileNotFoundError(path)
    return sorted(
        (p for p in path.rglob("*") if p.is_file() and p.suffix.lower() in suffixes),
        key=lambda p: str(p).casefold(),
    )


def read_gaze(path: Path) -> tuple[pd.DataFrame, dict[str, Any]]:
    source = Path(path).expanduser().resolve()
    suffix = source.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        with pd.ExcelFile(source) as excel:
            sheet_names = list(excel.sheet_names)
        preferred = ["眼动数据", "EyeData", "GazeData", "gaze_data"]
        candidates = preferred + [name for name in sheet_names if name not in preferred]
        errors: list[str] = []
        for sheet in candidates:
            if sheet not in sheet_names:
                continue
            try:
                header = pd.read_excel(source, sheet_name=sheet, nrows=0)
                dtype = _nanosecond_text_dtype(header.columns)
                frame = pd.read_excel(source, sheet_name=sheet, dtype=dtype)
                columns = resolve_columns(frame)
                return frame, {
                    "kind": "excel", "sheet": sheet, "columns": columns,
                    **_source_provenance(source),
                }
            except ValueError as exc:
                errors.append(f"{sheet}: {exc}")
        raise ValueError("No Excel sheet contains recognizable time, X, and Y columns; " + "; ".join(errors))
    delimiter, encoding = _detect_text_format(source)
    header = pd.read_csv(source, sep=delimiter, encoding=encoding, nrows=0)
    dtype = _nanosecond_text_dtype(header.columns)
    frame = pd.read_csv(source, sep=delimiter, encoding=encoding, dtype=dtype)
    return frame, {
        "kind": "text", "delimiter": delimiter, "encoding": encoding,
        "columns": resolve_columns(frame),
        **_source_provenance(source),
    }


def read_analysis_gaze(
    path: str | Path,
    *,
    time_unit: str | None = None,
) -> tuple[pd.DataFrame, dict[str, Any]]:
    """Read a complete gaze table and append canonical analysis fields.

    Source columns and row order are retained. Raw coordinates are selected
    from explicit raw fields or the legacy X/Y fields; display-smoothed and
    interpolated coordinate columns are never used as a fallback. A row is
    gaze-valid only when the source explicitly marks it accepted/valid.
    Numeric legacy time needs a unit from its header or ``time_unit``.
    """

    source = Path(path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    if time_unit is not None and time_unit not in {"seconds", "milliseconds", "microseconds"}:
        raise ValueError("time_unit must be seconds, milliseconds, or microseconds")

    frame, read_metadata, columns = _read_analysis_frame(source)
    manifest, manifest_path = _read_analysis_manifest(source)
    source_hash = _sha256_file(source)
    n = len(frame)
    x_column, y_column = columns["x_raw"], columns["y_raw"]
    if x_column is None or y_column is None:
        raise ValueError("No raw gaze X/Y columns were found; smoothed coordinates are not an analysis fallback")

    x = pd.to_numeric(frame[x_column], errors="coerce").to_numpy(dtype=float, na_value=np.nan)
    y = pd.to_numeric(frame[y_column], errors="coerce").to_numpy(dtype=float, na_value=np.nan)
    time_s, time_source, clock_quality, time_reason = _analysis_clock(
        frame, columns, manifest, time_unit
    )
    status = _analysis_status(frame, columns)
    interpolated = _analysis_flags(frame, columns["interpolated"])
    bad_segment = _analysis_flags(frame, columns["bad_segment"])
    explicit_valid_column = columns["gaze_valid"]
    accepted_column = columns["sample_accepted"]
    status_column = columns["sample_status"]
    valid_explicit = explicit_valid_column is not None or accepted_column is not None or status_column is not None
    if explicit_valid_column is not None:
        explicit_valid = _analysis_flags(frame, explicit_valid_column)
    elif accepted_column is not None:
        explicit_valid = _analysis_flags(frame, accepted_column)
    elif status_column is not None:
        explicit_valid = np.asarray([value == "accepted" for value in status], dtype=bool)
    else:
        explicit_valid = np.zeros(n, dtype=bool)

    # A non-accepted source status (for example blink/no_face) always keeps a
    # row out of gaze metrics even if another mask is contradictory.
    status_accepted = np.asarray([value == "accepted" for value in status], dtype=bool)
    if status_column is not None or accepted_column is not None or columns["blink"] is not None:
        explicit_valid &= status_accepted
    valid = explicit_valid & np.isfinite(x) & np.isfinite(y) & ~interpolated & ~bad_segment

    break_before = _analysis_flags(frame, columns["break_before"])
    if n:
        break_before[0] = True
    segment = _analysis_segment_ids(frame, columns, manifest)
    if n > 1:
        for i in range(1, n):
            if segment[i] != segment[i - 1]:
                break_before[i] = True

    reasons: list[str | None] = []
    raw_reason_column = columns["invalid_reason"]
    for i in range(n):
        row_reasons: list[str] = []
        if raw_reason_column is not None:
            raw_reason = frame.iloc[i][raw_reason_column]
            if not _analysis_missing(raw_reason):
                reason_text = str(raw_reason).strip()
                if reason_text and reason_text.casefold() not in {"none", "nan", "null"}:
                    row_reasons.append(reason_text)
        if not np.isfinite(x[i]) or not np.isfinite(y[i]):
            row_reasons.append("nonfinite_raw_coordinates")
        if not valid_explicit:
            row_reasons.append("missing_explicit_validity_status")
        elif not explicit_valid[i]:
            row_reasons.append("source_not_marked_valid")
        if status[i] != "accepted":
            row_reasons.append("sample_status_" + status[i])
        if interpolated[i]:
            row_reasons.append("interpolated_coordinate")
        if bad_segment[i]:
            row_reasons.append("bad_segment")
        if time_reason and not np.isfinite(time_s[i]):
            row_reasons.append(time_reason)
        reasons.append(";".join(dict.fromkeys(row_reasons)) or None)

    output = frame.copy()
    output["analysis_row_index"] = np.arange(n, dtype=np.int64)
    output["analysis_x_raw"] = x
    output["analysis_y_raw"] = y
    output["analysis_time_s"] = time_s
    output["analysis_time_source"] = time_source
    output["analysis_clock_quality"] = clock_quality
    output["analysis_sample_status"] = status
    output["analysis_gaze_valid"] = valid
    output["analysis_invalid_reason"] = reasons
    output["analysis_interpolated"] = interpolated
    output["analysis_bad_segment"] = bad_segment
    output["analysis_break_before"] = break_before
    output["analysis_segment_id"] = segment
    output["analysis_trial_id"] = _copy_analysis_identity(frame, columns["trial_id"])
    output["analysis_stimulus_id"] = _copy_analysis_identity(frame, columns["stimulus_id"])

    coordinate_reference = manifest.get("coordinate_reference")
    if not isinstance(coordinate_reference, dict):
        coordinate_reference = {"space": "unknown", "source": "unknown"}
    coordinate_status = manifest.get("coordinate_reference_status", "unknown")
    if coordinate_status != "verified":
        coordinate_status = "unknown" if coordinate_reference.get("space") == "unknown" else "declared_unverified"
    metadata = {
        **read_metadata,
        "source_path": str(source),
        "source_sha256": source_hash,
        "source_manifest_path": str(manifest_path) if manifest_path else None,
        "session_id": manifest.get("session_id", source.stem),
        "data_contract_version": manifest.get("data_contract_version"),
        "software_version": manifest.get("software_version", manifest.get("version")),
        "source_field_provenance": manifest.get("field_provenance", {}),
        "source_columns": [str(name) for name in frame.columns],
        "analysis_field_provenance": {
            "analysis_x_raw": {"source": "raw", "source_column": x_column},
            "analysis_y_raw": {"source": "raw", "source_column": y_column},
            "analysis_time_s": {"source": "derived", "derived_from": [time_source[0] if n else "unavailable"]},
            "analysis_gaze_valid": {
                "source": "derived",
                "derived_from": [
                    value for value in (explicit_valid_column, accepted_column, status_column, columns["interpolated"], columns["bad_segment"]) if value
                ],
            },
            "analysis_sample_status": {
                "source": "raw" if status_column is not None else "derived",
                "source_column": status_column,
                "derived_from": [value for value in (accepted_column, explicit_valid_column, columns["blink"]) if value] if status_column is None else [],
            },
            "analysis_break_before": {"source": "raw_and_normalized", "source_column": columns["break_before"]},
            "analysis_segment_id": {"source": "derived", "derived_from": [value for value in (columns["session_id"], columns["segment_id"]) if value]},
        },
        "coordinate_reference": coordinate_reference,
        "coordinate_reference_status": coordinate_status,
        "field_mapping": {
            "analysis_x_raw": x_column,
            "analysis_y_raw": y_column,
            "received_monotonic_ns": columns["received_monotonic_ns"],
            "elapsed_s": columns["elapsed_s"],
            "legacy_time": columns["time"],
            "analysis_gaze_valid": explicit_valid_column or accepted_column or status_column,
            "analysis_sample_status": status_column,
            "analysis_interpolated": columns["interpolated"],
            "analysis_bad_segment": columns["bad_segment"],
            "analysis_break_before": columns["break_before"],
            "analysis_segment_id": columns["segment_id"],
            "analysis_trial_id": columns["trial_id"],
            "analysis_stimulus_id": columns["stimulus_id"],
        },
        "time_source": time_source[0] if n else "unavailable",
        "clock_quality": sorted(set(clock_quality.tolist())) if n else ["unknown"],
        "time_unit_override": time_unit,
        "time_reason": time_reason,
        "analysis_preset": "raw_measured_full_contract_v1",
    }
    for source_metadata_field in ("model_feature_contract", "calibration_fit_residuals_px"):
        if source_metadata_field in manifest:
            metadata[source_metadata_field] = manifest[source_metadata_field]
    output.attrs["analysis_metadata"] = metadata
    return output, metadata


def _read_analysis_frame(path: Path) -> tuple[pd.DataFrame, dict[str, Any], dict[str, str | None]]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        with pd.ExcelFile(path) as excel:
            sheet_names = list(excel.sheet_names)
        preferred = ["眼动数据", "EyeData", "GazeData", "gaze_data"]
        candidates = preferred + [name for name in sheet_names if name not in preferred]
        errors: list[str] = []
        for sheet in candidates:
            if sheet not in sheet_names:
                continue
            try:
                header = pd.read_excel(path, sheet_name=sheet, nrows=0)
                columns = _analysis_column_map(header)
                if columns["x_raw"] is not None and columns["y_raw"] is not None:
                    # Preserve source ns text before pandas can infer a nullable
                    # large-integer column as float64 and round values above 2^53.
                    dtype = _nanosecond_text_dtype(header.columns)
                    frame = pd.read_excel(path, sheet_name=sheet, dtype=dtype)
                    return frame, {"kind": "excel", "sheet": sheet}, columns
            except (ValueError, KeyError) as exc:
                errors.append(f"{sheet}: {exc}")
        detail = "; ".join(errors)
        suffix = f" ({detail})" if detail else ""
        raise ValueError("No Excel sheet contains explicit raw gaze X/Y columns" + suffix)
    if suffix not in DATA_SUFFIXES:
        raise ValueError(f"Unsupported gaze input format: {path.suffix}")
    delimiter, encoding = _detect_text_format(path)
    header = pd.read_csv(path, sep=delimiter, encoding=encoding, nrows=0)
    columns = _analysis_column_map(header)
    if columns["x_raw"] is None or columns["y_raw"] is None:
        raise ValueError("No explicit raw gaze X/Y columns were found; smoothed coordinates are not an analysis fallback")
    # Pin every nanosecond-named field to object dtype before pandas can infer a
    # nullable large-integer column as float64 and round values above 2^53.
    dtype = _nanosecond_text_dtype(header.columns)
    frame = pd.read_csv(path, sep=delimiter, encoding=encoding, dtype=dtype)
    return frame, {"kind": "text", "delimiter": delimiter, "encoding": encoding}, columns


def _analysis_column_map(frame: pd.DataFrame) -> dict[str, str | None]:
    return {
        role: _pick_analysis_column(frame, aliases, role)
        for role, aliases in ANALYSIS_ALIASES.items()
    }


def _pick_analysis_column(frame: pd.DataFrame, aliases: tuple[str, ...], role: str) -> str | None:
    by_normalized: dict[str, list[str]] = {}
    for name in frame.columns:
        by_normalized.setdefault(_normalize(name), []).append(str(name))
    for alias in aliases:
        matches = by_normalized.get(_normalize(alias), [])
        if len(matches) > 1:
            raise ValueError(f"Ambiguous {role} columns normalize to the same name: {matches}")
        if matches:
            return matches[0]
    return None


def _read_analysis_manifest(source: Path) -> tuple[dict[str, Any], Path | None]:
    output_root = source.parent
    capture_manifest_name = f"{source.stem}.manifest.json"
    internal_manifest = sidecar_path(output_root, capture_manifest_name, layout=INTERNAL_LAYOUT)
    legacy_manifest = sidecar_path(output_root, capture_manifest_name, layout=LEGACY_LAYOUT)
    internal_exists = internal_manifest.exists() or internal_manifest.is_symlink()
    legacy_exists = legacy_manifest.exists() or legacy_manifest.is_symlink()
    if internal_exists and legacy_exists:
        raise ValueError(f"Conflicting root and .internal capture manifests exist: {capture_manifest_name}")
    if internal_exists or legacy_exists:
        candidate = internal_manifest if internal_exists else legacy_manifest
        if candidate.is_symlink() or not candidate.is_file():
            raise ValueError(f"Capture manifest is not a regular file: {candidate}")
        value = json.loads(candidate.read_text(encoding="utf-8"))
        if not isinstance(value, dict):
            raise ValueError(f"Analysis manifest must contain a JSON object: {candidate}")
        declared = normalize_sidecar_layout(value.get("sidecar_layout"))
        expected = existing_sidecar_path(output_root, capture_manifest_name, layout=declared)
        if declared != (INTERNAL_LAYOUT if internal_exists else LEGACY_LAYOUT) or expected != candidate:
            raise ValueError(f"Capture manifest path does not match its declared layout: {candidate}")
        files = value.get("files")
        recorded_manifest = files.get("manifest") if isinstance(files, dict) else None
        if recorded_manifest is not None:
            try:
                owned_sidecar_path(
                    output_root,
                    recorded_manifest,
                    name=capture_manifest_name,
                    layout=declared,
                )
            except (OSError, RuntimeError, TypeError, ValueError) as exc:
                raise ValueError(f"Capture manifest contains a conflicting path reference: {candidate}") from exc
        return value, candidate

    candidates = [source.with_suffix(".manifest.json"), source.with_suffix(source.suffix + ".manifest.json")]
    for candidate in candidates:
        if candidate.is_file():
            value = json.loads(candidate.read_text(encoding="utf-8"))
            if not isinstance(value, dict):
                raise ValueError(f"Analysis manifest must contain a JSON object: {candidate}")
            return value, candidate
    if source.suffix.lower() in {".xlsx", ".xlsm"}:
        try:
            table = pd.read_excel(source, sheet_name="CaptureMetadata")
        except (ValueError, KeyError):
            return {}, None
        if not {"Key", "Value"}.issubset(table.columns):
            return {}, None
        values: dict[str, Any] = {}
        for _, row in table.iterrows():
            key = row["Key"]
            if _analysis_missing(key):
                continue
            value = row["Value"]
            if isinstance(value, str):
                try:
                    value = json.loads(value)
                except json.JSONDecodeError:
                    pass
            values[str(key)] = value
        return values, None
    return {}, None


def _analysis_clock(
    frame: pd.DataFrame,
    columns: dict[str, str | None],
    manifest: dict[str, Any],
    time_unit: str | None,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, str | None]:
    n = len(frame)
    session_column = columns["session_id"]
    session_ids = _copy_analysis_identity(frame, session_column)
    if session_column is None:
        session_ids = np.full(n, str(manifest.get("session_id", "single_input_session")), dtype=object)
    else:
        for index, value in enumerate(session_ids):
            if value is None:
                session_ids[index] = f"unknown-session-row-{index}"

    candidates: list[tuple[str, str | None]] = [
        ("received_monotonic_ns", columns["received_monotonic_ns"]),
        ("elapsed_s", columns["elapsed_s"]),
        ("analysis_time_s", _pick_analysis_column(frame, ("analysis_time_s",), "normalized time")),
        ("legacy_time", columns["time"]),
    ]
    rejected_clocks: list[str] = []
    for source_name, column in candidates:
        if column is None:
            continue
        if source_name == "received_monotonic_ns":
            values, precision_lost = _monotonic_ns_values(frame[column])
            if precision_lost:
                times = np.full(n, np.nan)
                quality = np.full(n, "unreliable", dtype=object)
                reason = "monotonic_nanosecond_precision_lost"
            else:
                times, quality, reason = _relative_integer_clock(values, session_ids)
        elif source_name == "elapsed_s" or source_name == "analysis_time_s":
            values = pd.to_numeric(frame[column], errors="coerce").to_numpy(dtype=float, na_value=np.nan)
            times, quality, reason = _relative_clock(values, session_ids, 1.0, "session_elapsed")
            if source_name == "analysis_time_s":
                original_quality_column = _pick_analysis_column(frame, ANALYSIS_ALIASES["clock_quality"], "clock quality")
                if original_quality_column is not None:
                    original_quality = frame[original_quality_column].fillna("unknown").astype(str).to_numpy(dtype=object)
                    quality = np.asarray([
                        str(value).strip().casefold() if str(value).strip() else "unknown"
                        for value in original_quality
                    ], dtype=object)
        else:
            series = frame[column]
            if pd.api.types.is_datetime64_any_dtype(series.dtype):
                parsed = pd.to_datetime(series, errors="coerce")
                numeric_seconds = parsed.astype("int64", copy=False).to_numpy(dtype=float) / 1e9
                numeric_seconds[parsed.isna().to_numpy()] = np.nan
                times, quality, reason = _relative_clock(numeric_seconds, session_ids, 1.0, "legacy_datetime")
            else:
                numeric = pd.to_numeric(series, errors="coerce")
                nonmissing = series.notna()
                if nonmissing.any() and numeric[nonmissing].notna().all():
                    unit = _time_unit_from_header(column) or time_unit
                    if unit is None:
                        return (
                            np.full(n, np.nan),
                            np.full(n, "unknown", dtype=object),
                            np.full(n, "unknown", dtype=object),
                            "unknown_numeric_time_unit",
                        )
                    scale = {"seconds": 1.0, "milliseconds": 1e-3, "microseconds": 1e-6}[unit]
                    values = numeric.to_numpy(dtype=float, na_value=np.nan)
                    times, quality, reason = _relative_clock(values, session_ids, scale, "legacy_explicit_unit")
                else:
                    parsed = pd.to_datetime(series, errors="coerce")
                    if parsed.isna().all():
                        continue
                    numeric_seconds = parsed.astype("int64", copy=False).to_numpy(dtype=float) / 1e9
                    numeric_seconds[parsed.isna().to_numpy()] = np.nan
                    times, quality, reason = _relative_clock(numeric_seconds, session_ids, 1.0, "legacy_datetime")
        if not np.isfinite(times).any() and reason is None:
            continue
        if reason:
            rejected_clocks.append(f"{source_name}:{reason}")
            continue
        source_values = np.full(n, source_name, dtype=object)
        return times, source_values, quality, ";".join(rejected_clocks) or None

    return (
        np.full(n, np.nan),
        np.full(n, "unavailable", dtype=object),
        np.full(n, "unknown", dtype=object),
        ";".join(rejected_clocks) or "no_trusted_time_field",
    )


def _relative_clock(
    values: np.ndarray,
    groups: np.ndarray,
    scale: float,
    quality_name: str,
) -> tuple[np.ndarray, np.ndarray, str | None]:
    times = np.full(len(values), np.nan)
    quality = np.full(len(values), "missing", dtype=object)
    saw_nonmonotonic = False
    for group in dict.fromkeys(groups.tolist()):
        positions = np.flatnonzero(groups == group)
        local = values[positions]
        finite_positions = np.flatnonzero(np.isfinite(local))
        if not finite_positions.size:
            continue
        finite_values = local[finite_positions]
        if finite_values.size > 1 and np.any(np.diff(finite_values) < 0):
            quality[positions[finite_positions]] = "invalid"
            saw_nonmonotonic = True
            continue
        first = finite_values[0]
        times[positions[finite_positions]] = (finite_values - first) * scale
        quality[positions[finite_positions]] = quality_name
    return times, quality, "nonmonotonic_clock" if saw_nonmonotonic else None


def _monotonic_ns_values(series: pd.Series) -> tuple[np.ndarray, bool]:
    values = np.full(len(series), None, dtype=object)
    precision_lost = False
    for index, value in enumerate(series):
        if _analysis_missing(value) or isinstance(value, bool):
            continue
        if isinstance(value, (int, np.integer)):
            values[index] = int(value)
            continue
        if isinstance(value, (float, np.floating)) and math.isfinite(float(value)):
            if abs(float(value)) > 2**53:
                precision_lost = True
                continue
        try:
            parsed = Decimal(str(value).strip())
        except (InvalidOperation, ValueError):
            continue
        if parsed.is_finite() and parsed == parsed.to_integral_value():
            values[index] = int(parsed)
    return values, precision_lost


def _relative_integer_clock(
    values: np.ndarray,
    groups: np.ndarray,
) -> tuple[np.ndarray, np.ndarray, str | None]:
    times = np.full(len(values), np.nan)
    quality = np.full(len(values), "missing", dtype=object)
    saw_nonmonotonic = False
    for group in dict.fromkeys(groups.tolist()):
        positions = np.flatnonzero(groups == group)
        finite_positions = [position for position in positions if values[position] is not None]
        if not finite_positions:
            continue
        integer_values = [int(values[position]) for position in finite_positions]
        if any(right < left for left, right in zip(integer_values, integer_values[1:])):
            quality[finite_positions] = "invalid"
            saw_nonmonotonic = True
            continue
        origin = integer_values[0]
        for position, value in zip(finite_positions, integer_values):
            times[position] = (value - origin) / 1_000_000_000.0
            quality[position] = "measured_monotonic"
    return times, quality, "nonmonotonic_clock" if saw_nonmonotonic else None


def _analysis_status(frame: pd.DataFrame, columns: dict[str, str | None]) -> np.ndarray:
    status_column = columns["sample_status"]
    accepted_column = columns["sample_accepted"]
    valid_column = columns["gaze_valid"]
    blink_column = columns["blink"]
    if status_column is not None:
        result = []
        for value in frame[status_column]:
            status = "unknown" if _analysis_missing(value) else str(value).strip().casefold()
            result.append(status if status in ANALYSIS_STATUS_VALUES else status or "unknown")
        return np.asarray(result, dtype=object)
    if accepted_column is not None:
        result = np.where(_analysis_flags(frame, accepted_column), "accepted", "invalid").astype(object)
    elif valid_column is not None:
        result = np.where(_analysis_flags(frame, valid_column), "accepted", "invalid").astype(object)
    else:
        result = np.full(len(frame), "unknown", dtype=object)
    if blink_column is not None:
        blink = _analysis_flags(frame, blink_column)
        result[blink] = "blink"
    return result


def _analysis_segment_ids(
    frame: pd.DataFrame,
    columns: dict[str, str | None],
    manifest: dict[str, Any],
) -> np.ndarray:
    n = len(frame)
    session_column = columns["session_id"]
    if session_column is None:
        sessions = np.full(n, str(manifest.get("session_id", "single_input_session")), dtype=object)
    else:
        sessions = _copy_analysis_identity(frame, session_column)
        for index, value in enumerate(sessions):
            if value is None:
                sessions[index] = f"unknown-session-row-{index}"
    segment_column = columns["segment_id"]
    if segment_column is None:
        source_segments = np.full(n, None, dtype=object)
    else:
        source_segments = _copy_analysis_identity(frame, segment_column)
    return np.asarray([
        str(session) if segment is None else f"{session}::{segment}"
        for session, segment in zip(sessions, source_segments)
    ], dtype=object)


def _copy_analysis_identity(frame: pd.DataFrame, column: str | None) -> np.ndarray:
    if column is None:
        return np.full(len(frame), None, dtype=object)
    return np.asarray([
        None if _analysis_missing(value) else str(value)
        for value in frame[column]
    ], dtype=object)


def _analysis_flags(frame: pd.DataFrame, column: str | None) -> np.ndarray:
    if column is None:
        return np.zeros(len(frame), dtype=bool)

    def parse(value: Any) -> bool:
        if _analysis_missing(value):
            return False
        if isinstance(value, str):
            normalized = value.strip().casefold()
            return normalized in {"true", "1", "yes", "y"}
        return bool(value) if value in (True, False, 0, 1) else False

    return np.asarray([parse(value) for value in frame[column]], dtype=bool)


def _analysis_missing(value: Any) -> bool:
    if value is None:
        return True
    try:
        result = pd.isna(value)
        return bool(result) if np.isscalar(result) else False
    except (TypeError, ValueError):
        return False


def _time_unit_from_header(column: str) -> str | None:
    label = _normalize(column)
    if label.endswith("ms") or "millisecond" in label or "毫秒" in label:
        return "milliseconds"
    if label.endswith("us") or "microsecond" in label or "微秒" in label:
        return "microseconds"
    if label.endswith("s") or "seconds" in label or "秒" in label:
        return "seconds"
    return None


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _nanosecond_text_dtype(columns) -> dict[str, type] | None:
    nanosecond_columns = {
        str(column): object
        for column in columns
        if _is_nanosecond_column(str(column))
    }
    return nanosecond_columns or None


def _is_nanosecond_column(name: str) -> bool:
    normalized = _normalize(name)
    raw = name.casefold().strip()
    return (
        "monotonicns" in normalized
        or normalized == "receivedns"
        or raw.endswith("_ns")
        or raw.endswith(" ns")
    )


def _source_provenance(source: Path) -> dict[str, Any]:
    manifest, manifest_path = _read_analysis_manifest(source)
    return {
        "source_path": str(source),
        "source_sha256": _sha256_file(source),
        "source_manifest_path": str(manifest_path) if manifest_path is not None else None,
        "source_manifest_sha256": _sha256_file(manifest_path) if manifest_path is not None else None,
        "source_capture_manifest": manifest,
    }


def resolve_columns(frame: pd.DataFrame) -> dict[str, str | None]:
    normalized = {_normalize(name): str(name) for name in frame.columns}
    found: dict[str, str | None] = {}
    for role, aliases in ALIASES.items():
        found[role] = next((normalized[_normalize(alias)] for alias in aliases if _normalize(alias) in normalized), None)
    missing = [role for role in ("time", "x", "y") if found[role] is None]
    if missing:
        raise ValueError(f"Missing required column roles: {', '.join(missing)}")
    return found


def read_calibration_error(model_path: Path) -> tuple[float | None, str, Path | None]:
    workbook = model_path
    if model_path.suffix.lower() == ".pkl":
        workbook = model_path.with_suffix(".xlsx")
    if workbook.suffix.lower() not in {".xlsx", ".xlsm"} or not workbook.is_file():
        return None, "calibration_workbook_not_found", None
    try:
        with pd.ExcelFile(workbook) as excel:
            sheets = list(excel.sheet_names)
        additional = _case_sheet(sheets, "AdditionalCalibration")
        if additional:
            table = pd.read_excel(workbook, sheet_name=additional)
            column = _case_column(table.columns, "MeanErrorPerRound")
            if column:
                values = pd.to_numeric(table[column], errors="coerce").to_numpy(dtype=float)
                values = values[np.isfinite(values) & (values >= 0)]
                if values.size:
                    return float(values[-1]), "additional_calibration_final", workbook
        nine = _case_sheet(sheets, "NinePointCalibration")
        if nine:
            table = pd.read_excel(workbook, sheet_name=nine)
            column = _case_column(table.columns, "Error")
            if column:
                values = pd.to_numeric(table[column], errors="coerce").to_numpy(dtype=float)
                values = values[np.isfinite(values) & (values >= 0)]
                if values.size:
                    return float(values.mean()), "nine_point_mean_fallback", workbook
        return None, "no_usable_calibration_error", workbook
    except Exception as exc:  # retained in the batch report
        return None, f"calibration_read_error: {type(exc).__name__}: {exc}", workbook


def write_result(
    source: Path,
    destination: Path,
    processed: pd.DataFrame,
    metadata: dict[str, Any],
    qc: dict[str, Any],
    windows: pd.DataFrame,
    bad_segments: pd.DataFrame,
    suspect_segments: pd.DataFrame,
) -> None:
    sidecar = destination.with_suffix(destination.suffix + ".qc.json")
    existing = [path for path in (destination, sidecar) if path.exists() or path.is_symlink()]
    if existing:
        raise FileExistsError(f"Output already exists: {existing}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    if metadata["kind"] == "excel":
        shutil.copy2(source, destination)
        keep_vba = source.suffix.lower() == ".xlsm"
        workbook = load_workbook(destination, keep_vba=keep_vba)
        sheet = workbook[metadata["sheet"]]
        source_columns = list(pd.read_excel(source, sheet_name=metadata["sheet"], nrows=0).columns)
        appended = [name for name in processed.columns if name not in source_columns]
        start_column = sheet.max_column + 1
        header_fill = PatternFill("solid", fgColor="D9EAF7")
        for offset, name in enumerate(appended):
            cell = sheet.cell(row=1, column=start_column + offset, value=name)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            values = processed[name].tolist()
            for row, value in enumerate(values, start=2):
                if pd.isna(value):
                    value = None
                elif isinstance(value, np.generic):
                    value = value.item()
                sheet.cell(row=row, column=start_column + offset, value=value)
        _replace_sheet(workbook, "Pre_QC", _qc_table(qc))
        _replace_sheet(workbook, "Pre_Windows", windows)
        _replace_sheet(workbook, "Pre_BadSegments", bad_segments)
        _replace_sheet(workbook, "Pre_SuspectSegments", suspect_segments)
        workbook.save(destination)
    else:
        processed.to_csv(
            destination, sep=metadata["delimiter"], encoding=metadata["encoding"],
            index=False, lineterminator="\n",
        )
        sidecar.write_text(json.dumps(qc, ensure_ascii=False, indent=2), encoding="utf-8")


def _replace_sheet(workbook, title: str, frame: pd.DataFrame) -> None:
    if title in workbook.sheetnames:
        del workbook[title]
    sheet = workbook.create_sheet(title)
    for column, name in enumerate(frame.columns, start=1):
        cell = sheet.cell(row=1, column=column, value=str(name))
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row_index, row in enumerate(frame.itertuples(index=False, name=None), start=2):
        for column, value in enumerate(row, start=1):
            if pd.isna(value):
                value = None
            elif isinstance(value, (dict, list, tuple)):
                value = json.dumps(value, ensure_ascii=False)
            elif isinstance(value, np.generic):
                value = value.item()
            sheet.cell(row=row_index, column=column, value=value)
    sheet.freeze_panes = "A2"


def _qc_table(qc: dict[str, Any]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    def visit(prefix: str, value: Any) -> None:
        if isinstance(value, dict):
            for key, child in value.items():
                visit(f"{prefix}.{key}" if prefix else str(key), child)
        else:
            rows.append({"Key": prefix, "Value": value})
    visit("", qc)
    return pd.DataFrame(rows)


def _detect_text_format(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()[:65536]
    encoding = "utf-8-sig"
    try:
        sample = raw.decode(encoding)
    except UnicodeDecodeError:
        encoding = "gb18030"
        sample = raw.decode(encoding)
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
    except csv.Error:
        delimiter = "\t" if path.suffix.lower() in {".tsv", ".txt"} else ","
    return delimiter, encoding


def _normalize(value: Any) -> str:
    return "".join(ch for ch in str(value).casefold().strip() if ch not in " _-()[]（）")


def _case_sheet(sheets: list[str], requested: str) -> str | None:
    return next((name for name in sheets if name.casefold() == requested.casefold()), None)


def _case_column(columns, requested: str) -> str | None:
    return next((name for name in columns if str(name).casefold() == requested.casefold()), None)
